"""Operational Analytics upload rules, downloads, and unit-economics tests."""
import io

import pytest
from openpyxl import Workbook, load_workbook

FORECAST_CSV = (
    "ASIN,Short Name,Category,Daily Run Rate,PO Price,PO Value,Total Budget,"
    "CCOGS + Ads Budget,Review Budget\n"
    "B0OPER0001,Main Cable,Cables,10,100,1000,3000,3000,0\n"
    "B0OPER0002,Travel Cable,Cables,5,200,1000,1500,1400,100\n"
)

ACTUAL_CSV = (
    "ASIN,Total Order,PO PRICE,ADS+Cogs,OPA Payment\n"
    "B0OPER0001,11,100,105,0\n"
)
FORECAST_BYTES = FORECAST_CSV.encode()
ACTUAL_BYTES = ACTUAL_CSV.encode()


async def upload_forecast(client, headers, content=FORECAST_BYTES, name="forecast.csv"):
    return await client.post(
        "/api/v1/operational/forecasts",
        params={"forecast_month": "2026-04"},
        files={"file": (name, content, "text/csv")},
        headers=headers,
    )


async def upload_actual(
    client,
    headers,
    content=ACTUAL_BYTES,
    name="actual.csv",
    report_date="2026-04-01",
):
    return await client.post(
        "/api/v1/operational/actuals",
        params={"report_date": report_date},
        files={"file": (name, content, "text/csv")},
        headers=headers,
    )


async def add_target_asins(
    client,
    headers,
    upload_id,
    content,
    effective_from="2026-04-28",
    name="target-amendment.csv",
):
    return await client.post(
        f"/api/v1/operational/forecasts/{upload_id}/amendments",
        params={"effective_from": effective_from},
        files={"file": (name, content, "text/csv")},
        headers=headers,
    )


@pytest.mark.asyncio
async def test_forecast_actual_dashboard_and_unit_economics(client, auth_headers):
    forecast = await upload_forecast(client, auth_headers)
    assert forecast.status_code == 201, forecast.text
    assert forecast.json()["row_count"] == 2
    assert forecast.json()["units"] == 450  # (10 + 5) × 30 April days

    actual = await upload_actual(client, auth_headers)
    assert actual.status_code == 201, actual.text
    assert actual.json()["row_count"] == 1
    assert actual.json()["units"] == 11
    assert actual.json()["source_row_count"] == 1
    assert actual.json()["source_units"] == 11
    assert actual.json()["unmatched_units"] == 0
    assert actual.json()["unmatched_asins"] == []

    response = await client.get(
        "/api/v1/operational/dashboard",
        params={"date_from": "2026-04-01", "date_to": "2026-04-01"},
        headers=auth_headers,
    )
    assert response.status_code == 200, response.text
    dashboard = response.json()
    summary = dashboard["summary"]
    assert dashboard["covered_days"] == 1
    assert summary["planned_units"] == 15
    assert summary["actual_units"] == 11  # omitted B0OPER0002 contributes zero
    assert summary["planned_po_value"] == 2000  # sum of target DRR × target PO price
    assert summary["actual_po_value"] == 1100  # actual DRR × actual PO price
    assert summary["planned_spend"] == 150
    assert summary["actual_spend"] == 105
    assert summary["target_cac"] == 9.78  # target CCOGS + Ads / target orders
    assert summary["cac"] == 9.55  # ₹105 CCOGS + Ads / 11 actual orders

    asin = dashboard["categories"][0]["asins"][0]
    assert dashboard["categories"][0]["target_cac"] == 9.78
    assert dashboard["categories"][0]["cac"] == 9.55
    assert asin["asin"] == "B0OPER0001"
    assert asin["planned_units"] == 10
    assert asin["actual_units"] == 11
    assert asin["planned_po_value"] == 1000
    assert asin["actual_po_value"] == 1100
    assert asin["planned_spend"] == 100
    assert asin["target_cac"] == 10
    assert asin["cac"] == 9.55
    assert asin["volume_adjusted_budget"] == 110
    assert asin["adjusted_spend_variance"] == -5
    assert asin["status"] == "healthy"  # ₹105 is productive, not overspend
    assert asin["timeline"][0]["expected_units"] == 10
    assert asin["timeline"][0]["actual_units"] == 11
    assert dashboard["categories"][0]["timeline"][0]["expected_units"] == 15

    timeline = dashboard["timeline"][0]
    assert timeline["planned_po_value"] == 2000
    assert timeline["actual_po_value"] == 1100
    assert timeline["actual_ccogs_ads"] == 105
    assert timeline["cumulative_expected_units"] == 15
    assert timeline["cumulative_actual_units"] == 11


@pytest.mark.asyncio
async def test_spend_above_adjusted_budget_and_tolerance_is_flagged(client, auth_headers):
    await upload_forecast(client, auth_headers)
    overspend = ACTUAL_CSV.replace(",105,0", ",116,0").encode()
    response = await upload_actual(client, auth_headers, overspend)
    assert response.status_code == 201

    dashboard = (
        await client.get(
            "/api/v1/operational/dashboard",
            params={"date_from": "2026-04-01", "date_to": "2026-04-01"},
            headers=auth_headers,
        )
    ).json()
    asin = dashboard["categories"][0]["asins"][0]
    assert asin["volume_adjusted_budget"] == 110
    assert asin["actual_spend"] == 116
    assert asin["status"] == "overspend"  # ₹116 is > ₹110 + the confirmed 5% tolerance


@pytest.mark.asyncio
async def test_daily_actual_upload_accepts_negative_ccogs_ads(client, auth_headers):
    await upload_forecast(client, auth_headers)
    negative_ccogs_ads = ACTUAL_CSV.replace(",105,0", ",-25,0").encode()

    response = await upload_actual(client, auth_headers, negative_ccogs_ads)

    assert response.status_code == 201, response.text
    assert response.json()["spend"] == -25

    dashboard = (
        await client.get(
            "/api/v1/operational/dashboard",
            params={"date_from": "2026-04-01", "date_to": "2026-04-01"},
            headers=auth_headers,
        )
    ).json()
    asin = dashboard["categories"][0]["asins"][0]
    assert asin["spend_breakdown"]["ccogs_ads"]["actual"] == -25
    assert asin["actual_spend"] == -25


@pytest.mark.asyncio
async def test_one_file_per_period_and_forecast_foreign_key(client, auth_headers):
    no_forecast = await upload_actual(client, auth_headers)
    assert no_forecast.status_code == 409
    assert "target before daily actuals" in no_forecast.json()["detail"]

    await upload_forecast(client, auth_headers)
    duplicate_forecast = await upload_forecast(client, auth_headers)
    assert duplicate_forecast.status_code == 409

    await upload_actual(client, auth_headers)
    duplicate_actual = await upload_actual(client, auth_headers)
    assert duplicate_actual.status_code == 409
    assert "Delete it before uploading" in duplicate_actual.json()["detail"]


@pytest.mark.asyncio
async def test_day_28_target_amendment_preserves_actuals_and_starts_forward(
    client, auth_headers
):
    forecast = await upload_forecast(client, auth_headers)
    forecast_id = forecast.json()["upload_id"]
    day_27 = await upload_actual(
        client,
        auth_headers,
        report_date="2026-04-27",
    )
    assert day_27.status_code == 201, day_27.text

    amendment_csv = (
        "ASIN,Short Name,Category,Daily Run Rate,PO Price,PO Value,Total Budget,"
        "CCOGS + Ads Budget,Review Budget\n"
        "B0OPER0003,Launch Cable,Cables,4,150,600,300,250,50\n"
    ).encode()
    amendment = await add_target_asins(
        client, auth_headers, forecast_id, amendment_csv
    )
    assert amendment.status_code == 201, amendment.text
    assert amendment.json()["period"] == "2026-04-28"
    assert amendment.json()["row_count"] == 1
    assert amendment.json()["units"] == 12  # DRR 4 × Apr 28-30

    day_28_csv = (
        ACTUAL_CSV
        + "B0OPER0003,3,150,40,5\n"
    ).encode()
    day_28 = await upload_actual(
        client,
        auth_headers,
        day_28_csv,
        report_date="2026-04-28",
    )
    assert day_28.status_code == 201, day_28.text
    assert day_28.json()["row_count"] == 2

    dashboard = (
        await client.get(
            "/api/v1/operational/dashboard",
            params={"date_from": "2026-04-27", "date_to": "2026-04-28"},
            headers=auth_headers,
        )
    ).json()
    assert dashboard["summary"]["planned_units"] == 34
    assert dashboard["summary"]["actual_units"] == 25
    assert dashboard["timeline"][0]["date"] == "2026-04-27"
    assert dashboard["timeline"][0]["expected_units"] == 15
    assert dashboard["timeline"][1]["date"] == "2026-04-28"
    assert dashboard["timeline"][1]["expected_units"] == 19
    launched = next(
        asin
        for category in dashboard["categories"]
        for asin in category["asins"]
        if asin["asin"] == "B0OPER0003"
    )
    assert launched["planned_units"] == 4
    assert launched["actual_units"] == 3
    assert [point["date"] for point in launched["timeline"]] == ["2026-04-28"]

    forecasts = (
        await client.get("/api/v1/operational/forecasts", headers=auth_headers)
    ).json()
    assert forecasts[0]["row_count"] == 3
    assert forecasts[0]["planned_units"] == 462
    assert forecasts[0]["total_budget"] == 4800

    amendments = (
        await client.get(
            "/api/v1/operational/forecast-amendments", headers=auth_headers
        )
    ).json()
    assert len(amendments) == 1
    assert amendments[0]["effective_from"] == "2026-04-28"
    amendment_download = await client.get(
        f"/api/v1/operational/forecast-amendments/{amendments[0]['id']}/download",
        headers=auth_headers,
    )
    assert amendment_download.status_code == 200
    assert amendment_download.content == amendment_csv

    actuals = (
        await client.get("/api/v1/operational/actuals", headers=auth_headers)
    ).json()
    assert {row["report_date"] for row in actuals} == {"2026-04-27", "2026-04-28"}
    blocked_delete = await client.delete(
        f"/api/v1/operational/forecasts/{forecast_id}", headers=auth_headers
    )
    assert blocked_delete.status_code == 409

    deleted_amendment = await client.delete(
        f"/api/v1/operational/forecast-amendments/{amendments[0]['id']}",
        headers=auth_headers,
    )
    assert deleted_amendment.status_code == 204
    assert (
        await client.get(
            "/api/v1/operational/forecast-amendments", headers=auth_headers
        )
    ).json() == []
    assert (
        await client.get(
            f"/api/v1/operational/forecast-amendments/{amendments[0]['id']}/download",
            headers=auth_headers,
        )
    ).status_code == 404

    preserved_actuals = (
        await client.get("/api/v1/operational/actuals", headers=auth_headers)
    ).json()
    assert {row["report_date"] for row in preserved_actuals} == {
        "2026-04-27",
        "2026-04-28",
    }
    assert next(
        row for row in preserved_actuals if row["report_date"] == "2026-04-28"
    )["row_count"] == 2

    dashboard_after_delete = (
        await client.get(
            "/api/v1/operational/dashboard",
            params={"date_from": "2026-04-27", "date_to": "2026-04-28"},
            headers=auth_headers,
        )
    ).json()
    assert dashboard_after_delete["summary"]["planned_units"] == 30
    assert dashboard_after_delete["summary"]["actual_units"] == 22
    assert all(
        asin["asin"] != "B0OPER0003"
        for category in dashboard_after_delete["categories"]
        for asin in category["asins"]
    )

    readded = await add_target_asins(
        client,
        auth_headers,
        forecast_id,
        amendment_csv,
        effective_from="2026-04-29",
        name="replacement-amendment.csv",
    )
    assert readded.status_code == 201, readded.text
    assert readded.json()["units"] == 8
    dashboard_after_readd = (
        await client.get(
            "/api/v1/operational/dashboard",
            params={"date_from": "2026-04-28", "date_to": "2026-04-29"},
            headers=auth_headers,
        )
    ).json()
    readded_asin = next(
        asin
        for category in dashboard_after_readd["categories"]
        for asin in category["asins"]
        if asin["asin"] == "B0OPER0003"
    )
    assert readded_asin["planned_units"] == 4
    assert readded_asin["actual_units"] == 0
    assert [point["date"] for point in readded_asin["timeline"]] == ["2026-04-29"]


@pytest.mark.asyncio
async def test_target_amendment_rejects_existing_asins_and_reported_dates(
    client, auth_headers
):
    forecast_id = (await upload_forecast(client, auth_headers)).json()["upload_id"]
    existing_asin = (
        "ASIN,Short Name,Category,Daily Run Rate,PO Price,PO Value,Total Budget,"
        "CCOGS + Ads Budget,Review Budget\n"
        "B0OPER0001,Duplicate,Cables,1,100,100,100,100,0\n"
    ).encode()
    duplicate = await add_target_asins(
        client, auth_headers, forecast_id, existing_asin
    )
    assert duplicate.status_code == 409
    assert "already in the target" in duplicate.json()["detail"]

    await upload_actual(
        client,
        auth_headers,
        report_date="2026-04-27",
    )
    new_asin = existing_asin.replace(b"B0OPER0001,Duplicate", b"B0OPER0003,New")
    backdated = await add_target_asins(
        client,
        auth_headers,
        forecast_id,
        new_asin,
        effective_from="2026-04-27",
    )
    assert backdated.status_code == 409
    assert "after the latest uploaded actuals" in backdated.json()["detail"]

    wrong_month = await add_target_asins(
        client,
        auth_headers,
        forecast_id,
        new_asin,
        effective_from="2026-05-01",
    )
    assert wrong_month.status_code == 422
    assert "within April 2026" in wrong_month.json()["detail"]

    amendments = await client.get(
        "/api/v1/operational/forecast-amendments", headers=auth_headers
    )
    assert amendments.json() == []


@pytest.mark.asyncio
async def test_target_with_amendments_can_be_deleted_when_no_actuals(client, auth_headers):
    forecast_id = (await upload_forecast(client, auth_headers)).json()["upload_id"]
    amendment_csv = (
        "ASIN,Short Name,Category,Daily Run Rate,PO Price,PO Value,Total Budget,"
        "CCOGS + Ads Budget,Review Budget\n"
        "B0OPER0003,Launch Cable,Cables,4,150,600,300,250,50\n"
    ).encode()
    amendment = await add_target_asins(
        client, auth_headers, forecast_id, amendment_csv
    )
    assert amendment.status_code == 201, amendment.text

    deleted = await client.delete(
        f"/api/v1/operational/forecasts/{forecast_id}", headers=auth_headers
    )

    assert deleted.status_code == 204
    assert (
        await client.get(
            "/api/v1/operational/forecast-amendments", headers=auth_headers
        )
    ).json() == []


@pytest.mark.asyncio
async def test_unknown_actual_asin_is_ignored_and_budget_is_validated(client, auth_headers):
    bad_budget = FORECAST_CSV.replace(",1500,1400,100\n", ",1499,1400,100\n").encode()
    response = await upload_forecast(client, auth_headers, bad_budget)
    assert response.status_code == 422
    assert "Total Budget must equal" in response.json()["detail"]

    await upload_forecast(client, auth_headers)
    unknown_row = "B0UNKNOWN1,9,100,50,0\n"
    actuals_with_unknown = (ACTUAL_CSV + unknown_row).encode()
    response = await upload_actual(client, auth_headers, actuals_with_unknown)
    assert response.status_code == 201, response.text
    result = response.json()
    assert result["row_count"] == 1
    assert result["units"] == 11
    assert result["po_value"] == 1100
    assert result["spend"] == 105
    assert result["source_row_count"] == 2
    assert result["source_units"] == 20
    assert result["unmatched_units"] == 9
    assert result["unmatched_asins"] == [{"asin": "B0UNKNOWN1", "orders": 9}]

    uploads = (await client.get("/api/v1/operational/actuals", headers=auth_headers)).json()
    assert uploads[0]["source_row_count"] == 2
    assert uploads[0]["source_units"] == 20
    assert uploads[0]["unmatched_units"] == 9
    assert uploads[0]["unmatched_asins"] == [{"asin": "B0UNKNOWN1", "orders": 9}]


@pytest.mark.asyncio
async def test_actual_upload_requires_at_least_one_target_asin(client, auth_headers):
    await upload_forecast(client, auth_headers)
    unknown_only = (
        "ASIN,Total Order,PO PRICE,ADS+Cogs,OPA Payment\n"
        "B0UNKNOWN1,9,100,50,0\n"
    ).encode()

    response = await upload_actual(client, auth_headers, unknown_only)

    assert response.status_code == 422
    assert "no rows matching ASINs in the monthly target" in response.json()["detail"]
    assert "B0UNKNOWN1 (9.0 orders)" in response.json()["detail"]


@pytest.mark.asyncio
async def test_zero_target_product_can_contribute_actuals(client, auth_headers):
    target_with_zero_product = (
        FORECAST_CSV
        + "B0OPER0003,New Cable,Cables,0,150,0,0,0,0\n"
    ).encode()
    forecast = await upload_forecast(client, auth_headers, target_with_zero_product)
    assert forecast.status_code == 201, forecast.text

    actual_with_zero_target = (
        ACTUAL_CSV
        + "B0OPER0003,9,150,75,0\n"
    ).encode()
    actual = await upload_actual(client, auth_headers, actual_with_zero_target)
    assert actual.status_code == 201, actual.text
    assert actual.json()["row_count"] == 2
    assert actual.json()["units"] == 20
    assert actual.json()["po_value"] == 2450

    dashboard = (
        await client.get(
            "/api/v1/operational/dashboard",
            params={"date_from": "2026-04-01", "date_to": "2026-04-01"},
            headers=auth_headers,
        )
    ).json()
    zero_target = next(
        row
        for row in dashboard["categories"][0]["asins"]
        if row["asin"] == "B0OPER0003"
    )
    assert zero_target["planned_units"] == 0
    assert zero_target["actual_units"] == 9
    assert zero_target["actual_po_value"] == 1350
    assert zero_target["target_cac"] is None
    assert zero_target["cac"] == 8.33
    assert zero_target["status"] == "no_target"


@pytest.mark.asyncio
async def test_target_po_value_must_equal_daily_run_rate_times_po_price(client, auth_headers):
    incorrect_po_value = FORECAST_CSV.replace(",10,100,1000,", ",10,100,999,").encode()

    response = await upload_forecast(client, auth_headers, incorrect_po_value)

    assert response.status_code == 422
    assert "PO Value must equal Daily Run Rate multiplied by PO Price" in response.json()["detail"]


@pytest.mark.asyncio
async def test_target_po_value_preserves_validated_uploaded_precision(client, auth_headers):
    precise = FORECAST_CSV.replace(
        "B0OPER0001,Main Cable,Cables,10,100,1000,",
        "B0OPER0001,Main Cable,Cables,10.004,100,1000.4,",
    ).encode()
    response = await upload_forecast(client, auth_headers, precise)
    assert response.status_code == 201, response.text

    dashboard = (
        await client.get(
            "/api/v1/operational/dashboard",
            params={"date_from": "2026-04-01", "date_to": "2026-04-03"},
            headers=auth_headers,
        )
    ).json()
    assert dashboard["summary"]["planned_po_value"] == 6001.2
    assert dashboard["timeline"][0]["planned_po_value"] == 2000.4


@pytest.mark.asyncio
async def test_management_original_download_export_and_delete_order(client, auth_headers):
    forecast_id = (await upload_forecast(client, auth_headers)).json()["upload_id"]
    actual_id = (await upload_actual(client, auth_headers)).json()["upload_id"]

    forecasts = (await client.get("/api/v1/operational/forecasts", headers=auth_headers)).json()
    actuals = (await client.get("/api/v1/operational/actuals", headers=auth_headers)).json()
    assert len(forecasts) == 1
    assert len(actuals) == 1

    original = await client.get(
        f"/api/v1/operational/forecasts/{forecast_id}/download", headers=auth_headers
    )
    assert original.status_code == 200
    assert original.content == FORECAST_CSV.encode()

    export = await client.get(
        "/api/v1/operational/dashboard/export",
        params={"date_from": "2026-04-01", "date_to": "2026-04-01"},
        headers=auth_headers,
    )
    assert export.status_code == 200
    assert b"Grand Total" in export.content
    assert b"B0OPER0001" in export.content
    assert b"Target CAC (Target CCOGS + Ads / Target Order)" in export.content
    assert b"Actual CAC (CCOGS + Ads / Actual Order)" in export.content

    blocked = await client.delete(
        f"/api/v1/operational/forecasts/{forecast_id}", headers=auth_headers
    )
    assert blocked.status_code == 409
    assert "daily actual files" in blocked.json()["detail"]

    assert (
        await client.delete(f"/api/v1/operational/actuals/{actual_id}", headers=auth_headers)
    ).status_code == 204
    assert (
        await client.delete(f"/api/v1/operational/forecasts/{forecast_id}", headers=auth_headers)
    ).status_code == 204


@pytest.mark.asyncio
async def test_xlsx_templates_and_upload(client, auth_headers):
    template = await client.get(
        "/api/v1/operational/templates/forecast",
        params={"file_format": "xlsx"},
        headers=auth_headers,
    )
    assert template.status_code == 200
    assert "Monthly%20Target.xlsx" in template.headers["content-disposition"]
    workbook = load_workbook(io.BytesIO(template.content))
    sheet = workbook.active
    assert sheet.title == "Monthly Target"
    assert sheet["A1"].value == "ASIN"
    assert [cell.value for cell in sheet[1]] == [
        "ASIN",
        "Short Name",
        "Category",
        "Daily Run Rate",
        "PO Price",
        "PO Value",
        "Total Budget",
        "CCOGS + Ads Budget",
        "Review Budget",
    ]
    sheet.append(["B0XLSX001", "XLSX Cable", "Cables", 11, 100, 1100, 3000, 2500, 500])
    output = io.BytesIO()
    workbook.save(output)

    response = await upload_forecast(
        client,
        auth_headers,
        output.getvalue(),
        "forecast.xlsx",
    )
    assert response.status_code == 201, response.text
    assert response.json()["row_count"] == 1
    assert response.json()["units"] == 330  # April: uploaded rounded DRR 11 × 30


@pytest.mark.asyncio
async def test_actual_template_uses_consolidated_spend_columns(client, auth_headers):
    template = await client.get(
        "/api/v1/operational/templates/actual",
        params={"file_format": "xlsx"},
        headers=auth_headers,
    )

    assert template.status_code == 200
    workbook = load_workbook(io.BytesIO(template.content))
    assert [cell.value for cell in workbook.active[1]] == [
        "ASIN",
        "Total Orders",
        "PO PRICE",
        "ADS+Cogs",
        "OPA Payment",
    ]


@pytest.mark.asyncio
async def test_actual_xlsx_uses_only_mapped_columns_and_computes_total_spend(
    client, auth_headers
):
    await upload_forecast(client, auth_headers)
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "ASIN",
            "Short Name",
            "Total Orders",
            "PO PRICE",
            "ADS+Cogs",
            "OPA Payment",
            "Total Spend",
            "Unused Formula Result",
        ]
    )
    sheet.append(
        ["B0OPER0001", "Ignored", 11, 100, 105, 7, 999999, "#DIV/0!"]
    )
    output = io.BytesIO()
    workbook.save(output)

    response = await upload_actual(client, auth_headers, output.getvalue(), "actual.xlsx")

    assert response.status_code == 201, response.text
    assert response.json()["row_count"] == 1
    assert response.json()["units"] == 11
    assert response.json()["spend"] == 112


@pytest.mark.asyncio
async def test_operational_routes_require_auth(client):
    response = await client.get("/api/v1/operational/dashboard")
    assert response.status_code == 401

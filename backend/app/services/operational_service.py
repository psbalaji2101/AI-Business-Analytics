"""Operational plan/actual ingestion, management, and unit-economics analytics."""
import calendar
import csv
import hashlib
import io
import re
from collections import defaultdict
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.db.models import (
    OperationalActualRow,
    OperationalActualUpload,
    OperationalForecastAmendment,
    OperationalForecastRow,
    OperationalForecastUpload,
)
from app.schemas.operational import (
    ActualUploadOut,
    AsinOperationalMetrics,
    CategoryOperationalMetrics,
    ForecastAmendmentOut,
    ForecastUploadOut,
    OperationalActualUploadResult,
    OperationalDashboard,
    OperationalMetrics,
    OperationalTimelinePoint,
    OperationalUploadResult,
    SpendBreakdown,
    SpendComponent,
    UnmatchedActualAsin,
)

FORECAST_HEADERS = [
    "ASIN",
    "Short Name",
    "Category",
    "Daily Run Rate",
    "PO Price",
    "PO Value",
    "Total Budget",
    "CCOGS + Ads Budget",
    "Review Budget",
]

ACTUAL_HEADERS = [
    "ASIN",
    "Total Orders",
    "PO PRICE",
    "ADS+Cogs",
    "OPA Payment",
]

ALIASES = {
    "asin": "asin",
    "shortname": "short_name",
    "productname": "short_name",
    "category": "category",
    "catagory": "category",
    "dailyrunrate": "daily_run_rate",
    "expecteddrr": "daily_run_rate",
    "drr": "daily_run_rate",
    "drractual": "actual_drr",
    "actualdrr": "actual_drr",
    "actualdailyrunrate": "actual_drr",
    "totalorder": "actual_drr",
    "totalorders": "actual_drr",
    "poprice": "po_price",
    "povalue": "po_value",
    "ccogsadsbudget": "ccogs_ads",
    "ccogsandadsbudget": "ccogs_ads",
    "ccogsadsspend": "ccogs_ads",
    "ccogsandadsspend": "ccogs_ads",
    "adscogs": "ccogs_ads",
    "reviewsbudget": "reviews",
    "reviewbudget": "reviews",
    "reviewsspend": "reviews",
    "reviewspend": "reviews",
    "reviews": "reviews",
    "opapayment": "reviews",
    "totalbudget": "total_budget",
    "totalbudgettospend": "total_budget",
    "totalplannedspend": "total_budget",
    "totalspend": "total_spend",
    "actualspend": "total_spend",
}

FORECAST_REQUIRED = {
    "asin",
    "short_name",
    "category",
    "daily_run_rate",
    "po_price",
    "po_value",
    "ccogs_ads",
    "reviews",
    "total_budget",
}
ACTUAL_REQUIRED = {"asin", "actual_drr", "po_price", "ccogs_ads", "reviews"}
ACTUAL_RECONCILIATION_REQUIRED = {"asin", "actual_drr"}
COMPONENTS = ("ccogs_ads", "reviews")
TOLERANCE = 0.05


class OperationalError(Exception):
    def __init__(self, message: str, status_code: int = 422) -> None:
        super().__init__(message)
        self.message = message
        self.status_code = status_code


def _month_start(value: date) -> date:
    return value.replace(day=1)


def _month_end(value: date) -> date:
    return value.replace(day=calendar.monthrange(value.year, value.month)[1])


def _days_in_month(value: date) -> int:
    return calendar.monthrange(value.year, value.month)[1]


def _dates(start: date, end: date):
    current = start
    while current <= end:
        yield current
        current += timedelta(days=1)


def _norm_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").strip().lower())


def _content_type(filename: str) -> str:
    return (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        if Path(filename).suffix.lower() == ".xlsx"
        else "text/csv; charset=utf-8"
    )


def _read_rows(filename: str, content: bytes) -> list[dict[str, Any]]:
    suffix = Path(filename).suffix.lower()
    try:
        if suffix == ".csv":
            reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig")))
            return [dict(row) for row in reader]
        if suffix == ".xlsx":
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            sheet = workbook.active
            values = sheet.iter_rows(values_only=True)
            headers = next(values, None)
            if not headers:
                return []
            return [dict(zip(headers, row, strict=False)) for row in values]
    except (
        UnicodeDecodeError,
        csv.Error,
        OSError,
        ValueError,
        BadZipFile,
        InvalidFileException,
    ) as exc:
        raise OperationalError(f"Could not read {filename}: {exc}") from exc
    raise OperationalError("Only .csv and .xlsx files are supported.")


def _canonicalize(rows: list[dict[str, Any]], required: set[str], kind: str) -> list[dict[str, Any]]:
    if not rows:
        raise OperationalError(f"The {kind} file is empty.")
    header_map: dict[Any, str] = {}
    for header in rows[0]:
        canonical = ALIASES.get(_norm_header(header))
        if canonical in required:
            header_map[header] = canonical
    missing = sorted(required - set(header_map.values()))
    if missing:
        display = ", ".join(name.replace("_", " ").title() for name in missing)
        raise OperationalError(f"Missing required {kind} columns: {display}.")
    return [{header_map[key]: value for key, value in row.items() if key in header_map} for row in rows]


def _number(
    value: Any, row_number: int, field: str, *, allow_negative: bool = False
) -> Decimal:
    if value is None or str(value).strip() == "":
        raise OperationalError(f"Row {row_number}: {field} is required.")
    cleaned = str(value).replace(",", "").replace("₹", "").strip()
    try:
        number = Decimal(cleaned)
    except InvalidOperation as exc:
        raise OperationalError(f"Row {row_number}: {field} must be numeric.") from exc
    if not number.is_finite() or (number < 0 and not allow_negative):
        raise OperationalError(f"Row {row_number}: {field} must be zero or greater.")
    return number


def _text(value: Any, row_number: int, field: str) -> str:
    result = str(value or "").strip()
    if not result:
        raise OperationalError(f"Row {row_number}: {field} is required.")
    return result


def _round(value: float | Decimal) -> float:
    return round(float(value), 2)


def _actual_reconciliation(
    raw_rows: list[dict[str, Any]], matched_asins: set[str]
) -> tuple[int, Decimal, dict[str, Decimal]]:
    """Summarize every source row and group ASINs excluded from analytics."""
    source_row_count = 0
    source_units = Decimal("0")
    unmatched: dict[str, Decimal] = {}

    for row_number, row in enumerate(raw_rows, start=2):
        if all(value is None or str(value).strip() == "" for value in row.values()):
            continue
        asin = _text(row.get("asin"), row_number, "ASIN").upper()
        orders = _number(row.get("actual_drr"), row_number, "Total Orders")
        source_row_count += 1
        source_units += orders
        if asin not in matched_asins:
            unmatched[asin] = unmatched.get(asin, Decimal("0")) + orders

    return source_row_count, source_units, unmatched


def _unmatched_actual_asins(values: dict[str, Decimal]) -> list[UnmatchedActualAsin]:
    return [
        UnmatchedActualAsin(asin=asin, orders=_round(orders))
        for asin, orders in values.items()
    ]


def _parse_target_rows(filename: str, content: bytes) -> list[dict[str, Any]]:
    raw_rows = _canonicalize(_read_rows(filename, content), FORECAST_REQUIRED, "target")
    parsed: list[dict[str, Any]] = []
    seen: set[str] = set()
    for row_number, row in enumerate(raw_rows, start=2):
        if all(value is None or str(value).strip() == "" for value in row.values()):
            continue
        asin = _text(row.get("asin"), row_number, "ASIN").upper()
        if len(asin) > 20:
            raise OperationalError(f"Row {row_number}: ASIN must be 20 characters or fewer.")
        if asin in seen:
            raise OperationalError(f"Row {row_number}: duplicate ASIN {asin} in target file.")
        seen.add(asin)
        item = {
            "asin": asin,
            "short_name": _text(row.get("short_name"), row_number, "Short Name"),
            "category": _text(row.get("category"), row_number, "Category"),
            "daily_run_rate": _number(row.get("daily_run_rate"), row_number, "Daily Run Rate"),
            "po_price": _number(row.get("po_price"), row_number, "PO Price"),
            "po_value": _number(row.get("po_value"), row_number, "PO Value"),
            "total_budget": _number(row.get("total_budget"), row_number, "Total Budget"),
            "ccogs_ads": _number(row.get("ccogs_ads"), row_number, "CCOGS + Ads Budget"),
            "reviews": _number(row.get("reviews"), row_number, "Review Budget"),
        }
        calculated_po_value = item["daily_run_rate"] * item["po_price"]
        if abs(calculated_po_value - item["po_value"]) > Decimal("0.01"):
            raise OperationalError(
                f"Row {row_number}: PO Value must equal Daily Run Rate multiplied by "
                f"PO Price (expected {_round(calculated_po_value)}, "
                f"found {_round(item['po_value'])})."
            )
        component_total = sum(item[name] for name in COMPONENTS)
        if abs(component_total - item["total_budget"]) > Decimal("0.01"):
            raise OperationalError(
                f'Row {row_number}: Total Budget must equal "CCOGS + Ads Budget" '
                '+ "Review Budget" '
                f"(expected {_round(component_total)}, found {_round(item['total_budget'])})."
            )
        parsed.append(item)

    if not parsed:
        raise OperationalError("The target file has no usable data rows.")
    return parsed


def _forecast_component(row: OperationalForecastRow, component: str) -> float:
    if component == "ccogs_ads":
        # Include legacy coupon values so previously uploaded files retain their total spend.
        return sum(
            float(value) for value in (row.ccogs_budget, row.ads_budget, row.coupons_budget)
        )
    return float(row.reviews_budget)


def _actual_component(row: OperationalActualRow, component: str) -> float:
    if component == "ccogs_ads":
        # New consolidated values are stored in ccogs_spend; the sum keeps old uploads compatible.
        return sum(float(value) for value in (row.ccogs_spend, row.ads_spend, row.coupons_spend))
    return float(row.reviews_spend)


def _blank_accumulator() -> dict[str, Any]:
    return {
        "planned_units": 0.0,
        "actual_units": 0.0,
        "planned_po_value": 0.0,
        "actual_po_value": 0.0,
        "planned": {component: 0.0 for component in COMPONENTS},
        "actual": {component: 0.0 for component in COMPONENTS},
        "adjusted": {component: 0.0 for component in COMPONENTS},
    }


def _merge(target: dict[str, Any], source: dict[str, Any]) -> None:
    for field in ("planned_units", "actual_units", "planned_po_value", "actual_po_value"):
        target[field] += source[field]
    for bucket in ("planned", "actual", "adjusted"):
        for component in COMPONENTS:
            target[bucket][component] += source[bucket][component]


def _metrics(values: dict[str, Any]) -> OperationalMetrics:
    planned_units = values["planned_units"]
    actual_units = values["actual_units"]
    planned_po = values["planned_po_value"]
    actual_po = values["actual_po_value"]
    planned_spend = sum(values["planned"].values())
    actual_spend = sum(values["actual"].values())
    adjusted_budget = sum(values["adjusted"].values())
    achievement = actual_units / planned_units * 100 if planned_units else None
    planned_cpu = planned_spend / planned_units if planned_units else None
    actual_cpu = actual_spend / actual_units if actual_units else None
    target_cac = values["planned"]["ccogs_ads"] / planned_units if planned_units else None
    cac = values["actual"]["ccogs_ads"] / actual_units if actual_units else None
    contribution = actual_po - actual_spend

    if planned_units <= 0:
        status = "no_target"
    elif actual_units <= 0:
        status = "overspend" if actual_spend > 0 else "no_sales"
    elif actual_spend <= adjusted_budget * (1 + TOLERANCE):
        status = "healthy" if (achievement or 0) >= 95 else "efficient_but_behind"
    else:
        status = "overspend"

    breakdown = {}
    for component in COMPONENTS:
        adjusted = values["adjusted"][component]
        actual = values["actual"][component]
        breakdown[component] = SpendComponent(
            planned=_round(values["planned"][component]),
            actual=_round(actual),
            adjusted_budget=_round(adjusted),
            adjusted_variance=_round(actual - adjusted),
        )

    return OperationalMetrics(
        planned_units=_round(planned_units),
        actual_units=_round(actual_units),
        unit_variance=_round(actual_units - planned_units),
        unit_achievement_pct=_round(achievement) if achievement is not None else None,
        planned_po_value=_round(planned_po),
        actual_po_value=_round(actual_po),
        po_value_variance=_round(actual_po - planned_po),
        planned_spend=_round(planned_spend),
        actual_spend=_round(actual_spend),
        nominal_spend_variance=_round(actual_spend - planned_spend),
        volume_adjusted_budget=_round(adjusted_budget),
        adjusted_spend_variance=_round(actual_spend - adjusted_budget),
        planned_cost_per_unit=_round(planned_cpu) if planned_cpu is not None else None,
        actual_cost_per_unit=_round(actual_cpu) if actual_cpu is not None else None,
        target_cac=_round(target_cac) if target_cac is not None else None,
        cac=_round(cac) if cac is not None else None,
        planned_spend_utilization_pct=_round(planned_spend / planned_po * 100)
        if planned_po
        else None,
        actual_spend_utilization_pct=_round(actual_spend / actual_po * 100)
        if actual_po
        else None,
        contribution_value=_round(contribution),
        contribution_margin_pct=_round(contribution / actual_po * 100) if actual_po else None,
        status=status,
        spend_breakdown=SpendBreakdown(**breakdown),
    )


def _build_timeline(
    plan_by_date: dict[date, list[OperationalForecastRow]],
    actual_by_row_date: dict[tuple[int, date], OperationalActualRow],
) -> list[OperationalTimelinePoint]:
    timeline: list[OperationalTimelinePoint] = []
    cumulative = {
        "expected_units": 0.0,
        "actual_units": 0.0,
        "planned_spend": 0.0,
        "actual_spend": 0.0,
        "adjusted_budget": 0.0,
    }
    for report_date in sorted(plan_by_date):
        expected_units = actual_units = planned_spend = actual_spend = adjusted = 0.0
        planned_po_value = actual_po_value = actual_ccogs_ads = 0.0
        for row in plan_by_date[report_date]:
            month_days = _days_in_month(report_date)
            expected_units += float(row.daily_run_rate)
            planned_po_value += float(row.po_value)
            planned_spend += float(row.total_budget) / month_days
            actual = actual_by_row_date.get((row.id, report_date))
            if not actual:
                continue
            units = float(actual.actual_drr)
            actual_units += units
            actual_po_value += units * float(actual.po_price)
            actual_ccogs_ads += _actual_component(actual, "ccogs_ads")
            actual_spend += float(actual.total_spend)
            monthly_units = float(row.daily_run_rate) * month_days
            if monthly_units:
                adjusted += float(row.total_budget) / monthly_units * units
        cumulative["expected_units"] += expected_units
        cumulative["actual_units"] += actual_units
        cumulative["planned_spend"] += planned_spend
        cumulative["actual_spend"] += actual_spend
        cumulative["adjusted_budget"] += adjusted
        timeline.append(
            OperationalTimelinePoint(
                date=report_date,
                expected_units=_round(expected_units),
                actual_units=_round(actual_units),
                planned_po_value=_round(planned_po_value),
                actual_po_value=_round(actual_po_value),
                actual_ccogs_ads=_round(actual_ccogs_ads),
                cumulative_expected_units=_round(cumulative["expected_units"]),
                cumulative_actual_units=_round(cumulative["actual_units"]),
                planned_spend=_round(planned_spend),
                actual_spend=_round(actual_spend),
                cumulative_planned_spend=_round(cumulative["planned_spend"]),
                cumulative_actual_spend=_round(cumulative["actual_spend"]),
                cumulative_adjusted_budget=_round(cumulative["adjusted_budget"]),
            )
        )
    return timeline


def _empty_dashboard() -> OperationalDashboard:
    return OperationalDashboard(
        date_from=None,
        date_to=None,
        covered_days=0,
        summary=_metrics(_blank_accumulator()),
        categories=[],
        timeline=[],
    )


class OperationalService:
    def __init__(self, db: AsyncSession) -> None:
        self.db = db

    async def upload_forecast(
        self, forecast_month: date, filename: str, content: bytes
    ) -> OperationalUploadResult:
        month = _month_start(forecast_month)
        existing = await self.db.scalar(
            select(OperationalForecastUpload).where(
                OperationalForecastUpload.forecast_month == month
            )
        )
        if existing:
            raise OperationalError(
                f"A target already exists for {month:%B %Y}. "
                "Use Add ASINs for new launches. A full replacement requires deleting the "
                "existing target first.",
                409,
            )

        parsed = _parse_target_rows(filename, content)
        month_days = _days_in_month(month)

        upload = OperationalForecastUpload(
            forecast_month=month,
            filename=filename,
            content_type=_content_type(filename),
            original_content=content,
            content_hash=hashlib.sha256(content).hexdigest(),
            row_count=len(parsed),
            planned_units=sum(float(row["daily_run_rate"]) * month_days for row in parsed),
            po_value=sum(float(row["po_value"]) for row in parsed),
            total_budget=sum(float(row["total_budget"]) for row in parsed),
        )
        self.db.add(upload)
        await self.db.flush()
        for row in parsed:
            self.db.add(
                OperationalForecastRow(
                    upload_id=upload.id,
                    asin=row["asin"],
                    short_name=row["short_name"],
                    category=row["category"],
                    effective_from=month,
                    daily_run_rate=row["daily_run_rate"],
                    po_price=row["po_price"],
                    po_value=row["po_value"],
                    ccogs_budget=row["ccogs_ads"],
                    ads_budget=0,
                    coupons_budget=0,
                    reviews_budget=row["reviews"],
                    total_budget=row["total_budget"],
                )
            )
        await self.db.commit()
        await self.db.refresh(upload)
        return OperationalUploadResult(
            upload_id=upload.id,
            filename=filename,
            row_count=upload.row_count,
            period=month,
            units=_round(upload.planned_units),
            po_value=_round(upload.po_value),
            spend=_round(upload.total_budget),
        )

    async def add_forecast_asins(
        self,
        upload_id: int,
        effective_from: date,
        filename: str,
        content: bytes,
    ) -> OperationalUploadResult:
        forecast = await self.db.get(OperationalForecastUpload, upload_id)
        if not forecast:
            raise OperationalError("Target upload not found.", 404)
        month = forecast.forecast_month
        if _month_start(effective_from) != month:
            raise OperationalError(
                f"Effective date must be within {month:%B %Y}."
            )

        latest_actual = await self.db.scalar(
            select(func.max(OperationalActualUpload.report_date)).where(
                OperationalActualUpload.forecast_upload_id == upload_id
            )
        )
        if latest_actual and effective_from <= latest_actual:
            next_date = latest_actual + timedelta(days=1)
            raise OperationalError(
                f"Effective date must be after the latest uploaded actuals "
                f"({latest_actual:%d %B %Y}). Choose {next_date.isoformat()} or later.",
                409,
            )

        parsed = _parse_target_rows(filename, content)
        incoming_asins = {row["asin"] for row in parsed}
        existing_rows = list(
            await self.db.scalars(
                select(OperationalForecastRow).where(
                    OperationalForecastRow.upload_id == upload_id,
                    OperationalForecastRow.asin.in_(incoming_asins),
                )
            )
        )
        source_amendment_ids = {
            row.source_amendment_id
            for row in existing_rows
            if row.source_amendment_id is not None
        }
        deleted_amendment_ids = (
            set(
                await self.db.scalars(
                    select(OperationalForecastAmendment.id).where(
                        OperationalForecastAmendment.id.in_(source_amendment_ids),
                        OperationalForecastAmendment.deleted_at.is_not(None),
                    )
                )
            )
            if source_amendment_ids
            else set()
        )
        active_existing_asins = {
            row.asin
            for row in existing_rows
            if row.source_amendment_id is None
            or row.source_amendment_id not in deleted_amendment_ids
        }
        if active_existing_asins:
            examples = ", ".join(sorted(active_existing_asins)[:10])
            remaining = len(active_existing_asins) - 10
            suffix = f" and {remaining} more" if remaining > 0 else ""
            raise OperationalError(
                f"This amendment contains {len(active_existing_asins)} ASIN(s) already in the "
                "target: "
                f"{examples}{suffix}. Upload only newly launched ASINs.",
                409,
            )
        inactive_rows_by_asin = {
            row.asin: row
            for row in existing_rows
            if row.source_amendment_id in deleted_amendment_ids
        }

        active_days = (_month_end(month) - effective_from).days + 1
        planned_units = sum(float(row["daily_run_rate"]) * active_days for row in parsed)
        po_value = sum(float(row["po_value"]) for row in parsed)
        total_budget = sum(float(row["total_budget"]) for row in parsed)
        amendment = OperationalForecastAmendment(
            forecast_upload_id=upload_id,
            effective_from=effective_from,
            filename=filename,
            content_type=_content_type(filename),
            original_content=content,
            content_hash=hashlib.sha256(content).hexdigest(),
            row_count=len(parsed),
            planned_units=planned_units,
            po_value=po_value,
            total_budget=total_budget,
        )
        self.db.add(amendment)
        await self.db.flush()
        for row in parsed:
            target_row = inactive_rows_by_asin.get(row["asin"])
            if target_row is None:
                target_row = OperationalForecastRow(upload_id=upload_id, asin=row["asin"])
                self.db.add(target_row)
            target_row.source_amendment_id = amendment.id
            target_row.effective_from = effective_from
            target_row.short_name = row["short_name"]
            target_row.category = row["category"]
            target_row.daily_run_rate = row["daily_run_rate"]
            target_row.po_price = row["po_price"]
            target_row.po_value = row["po_value"]
            target_row.ccogs_budget = row["ccogs_ads"]
            target_row.ads_budget = 0
            target_row.coupons_budget = 0
            target_row.reviews_budget = row["reviews"]
            target_row.total_budget = row["total_budget"]

        forecast.row_count = int(forecast.row_count) + len(parsed)
        forecast.planned_units = float(forecast.planned_units) + planned_units
        forecast.po_value = float(forecast.po_value) + po_value
        forecast.total_budget = float(forecast.total_budget) + total_budget
        await self.db.commit()
        await self.db.refresh(amendment)
        return OperationalUploadResult(
            upload_id=amendment.id,
            filename=filename,
            row_count=amendment.row_count,
            period=effective_from,
            units=_round(amendment.planned_units),
            po_value=_round(amendment.po_value),
            spend=_round(amendment.total_budget),
        )

    async def upload_actual(
        self, report_date: date, filename: str, content: bytes
    ) -> OperationalActualUploadResult:
        existing = await self.db.scalar(
            select(OperationalActualUpload).where(
                OperationalActualUpload.report_date == report_date
            )
        )
        if existing:
            raise OperationalError(
                f"An actuals file already exists for {report_date:%d %B %Y}. "
                "Delete it before uploading a replacement.",
                409,
            )
        forecast = await self.db.scalar(
            select(OperationalForecastUpload).where(
                OperationalForecastUpload.forecast_month == _month_start(report_date)
            )
        )
        if not forecast:
            raise OperationalError(
                f"Upload the {_month_start(report_date):%B %Y} target before daily actuals.", 409
            )
        forecast_rows = list(
            await self.db.scalars(
                select(OperationalForecastRow)
                .outerjoin(
                    OperationalForecastAmendment,
                    OperationalForecastRow.source_amendment_id
                    == OperationalForecastAmendment.id,
                )
                .where(
                    OperationalForecastRow.upload_id == forecast.id,
                    or_(
                        OperationalForecastRow.source_amendment_id.is_(None),
                        OperationalForecastAmendment.deleted_at.is_(None),
                    ),
                )
            )
        )
        by_asin = {
            row.asin: row
            for row in forecast_rows
            if (row.effective_from or forecast.forecast_month) <= report_date
        }
        raw_rows = _canonicalize(_read_rows(filename, content), ACTUAL_REQUIRED, "actuals")
        source_row_count, source_units, unmatched = _actual_reconciliation(
            raw_rows, set(by_asin)
        )
        parsed: list[dict[str, Any]] = []
        seen: set[str] = set()
        for row_number, row in enumerate(raw_rows, start=2):
            if all(value is None or str(value).strip() == "" for value in row.values()):
                continue
            asin = _text(row.get("asin"), row_number, "ASIN").upper()
            forecast_row = by_asin.get(asin)
            if not forecast_row:
                continue
            if asin in seen:
                raise OperationalError(f"Row {row_number}: duplicate ASIN {asin} in actuals file.")
            seen.add(asin)
            item = {
                "forecast_row": forecast_row,
                "actual_drr": _number(row.get("actual_drr"), row_number, "Total Orders"),
                "po_price": _number(row.get("po_price"), row_number, "PO Price"),
                "ccogs_ads": _number(
                    row.get("ccogs_ads"),
                    row_number,
                    "ADS+Cogs",
                    allow_negative=True,
                ),
                "reviews": _number(row.get("reviews"), row_number, "OPA Payment"),
            }
            item["total_spend"] = sum(item[name] for name in COMPONENTS)
            item["po_value"] = item["actual_drr"] * item["po_price"]
            parsed.append(item)
        if not parsed:
            unmatched_detail = ", ".join(
                f"{asin} ({_round(orders)} orders)" for asin, orders in unmatched.items()
            )
            suffix = f" Missing from target: {unmatched_detail}." if unmatched_detail else ""
            raise OperationalError(
                "The actuals file has no rows matching ASINs in the monthly target." + suffix
            )

        upload = OperationalActualUpload(
            forecast_upload_id=forecast.id,
            report_date=report_date,
            filename=filename,
            content_type=_content_type(filename),
            original_content=content,
            content_hash=hashlib.sha256(content).hexdigest(),
            row_count=len(parsed),
            actual_units=sum(float(row["actual_drr"]) for row in parsed),
            po_value=sum(float(row["po_value"]) for row in parsed),
            total_spend=sum(float(row["total_spend"]) for row in parsed),
        )
        self.db.add(upload)
        await self.db.flush()
        for row in parsed:
            self.db.add(
                OperationalActualRow(
                    upload_id=upload.id,
                    forecast_row_id=row["forecast_row"].id,
                    actual_drr=row["actual_drr"],
                    po_price=row["po_price"],
                    ccogs_spend=row["ccogs_ads"],
                    ads_spend=0,
                    coupons_spend=0,
                    reviews_spend=row["reviews"],
                    total_spend=row["total_spend"],
                    po_value=row["po_value"],
                )
            )
        await self.db.commit()
        await self.db.refresh(upload)
        unmatched_units = sum(unmatched.values(), Decimal("0"))
        return OperationalActualUploadResult(
            upload_id=upload.id,
            filename=filename,
            row_count=upload.row_count,
            period=report_date,
            units=_round(upload.actual_units),
            po_value=_round(upload.po_value),
            spend=_round(upload.total_spend),
            source_row_count=source_row_count,
            source_units=_round(source_units),
            unmatched_units=_round(unmatched_units),
            unmatched_asins=_unmatched_actual_asins(unmatched),
        )

    async def list_forecasts(self) -> list[ForecastUploadOut]:
        rows = await self.db.scalars(
            select(OperationalForecastUpload).order_by(
                OperationalForecastUpload.forecast_month.desc()
            )
        )
        return [ForecastUploadOut.model_validate(row) for row in rows]

    async def list_forecast_amendments(self) -> list[ForecastAmendmentOut]:
        rows = await self.db.scalars(
            select(OperationalForecastAmendment)
            .where(OperationalForecastAmendment.deleted_at.is_(None))
            .order_by(
                OperationalForecastAmendment.effective_from.desc(),
                OperationalForecastAmendment.uploaded_at.desc(),
            )
        )
        return [ForecastAmendmentOut.model_validate(row) for row in rows]

    async def list_actuals(self) -> list[ActualUploadOut]:
        rows = list(
            await self.db.scalars(
                select(OperationalActualUpload).order_by(
                    OperationalActualUpload.report_date.desc()
                )
            )
        )
        matches = await self.db.execute(
            select(OperationalActualRow.upload_id, OperationalForecastRow.asin).join(
                OperationalForecastRow,
                OperationalActualRow.forecast_row_id == OperationalForecastRow.id,
            )
        )
        matched_by_upload: dict[int, set[str]] = defaultdict(set)
        for upload_id, asin in matches:
            matched_by_upload[upload_id].add(asin)

        result: list[ActualUploadOut] = []
        for row in rows:
            raw_rows = _canonicalize(
                _read_rows(row.filename, row.original_content),
                ACTUAL_RECONCILIATION_REQUIRED,
                "actuals",
            )
            source_row_count, source_units, unmatched = _actual_reconciliation(
                raw_rows, matched_by_upload[row.id]
            )
            unmatched_units = sum(unmatched.values(), Decimal("0"))
            result.append(
                ActualUploadOut.model_validate(row).model_copy(
                    update={
                        "source_row_count": source_row_count,
                        "source_units": _round(source_units),
                        "unmatched_units": _round(unmatched_units),
                        "unmatched_asins": _unmatched_actual_asins(unmatched),
                    }
                )
            )
        return result

    async def original_forecast(self, upload_id: int) -> OperationalForecastUpload:
        upload = await self.db.get(OperationalForecastUpload, upload_id)
        if not upload:
            raise OperationalError("Target upload not found.", 404)
        return upload

    async def original_forecast_amendment(
        self, amendment_id: int
    ) -> OperationalForecastAmendment:
        amendment = await self.db.get(OperationalForecastAmendment, amendment_id)
        if not amendment or amendment.deleted_at is not None:
            raise OperationalError("Target amendment not found.", 404)
        return amendment

    async def original_actual(self, upload_id: int) -> OperationalActualUpload:
        upload = await self.db.get(OperationalActualUpload, upload_id)
        if not upload:
            raise OperationalError("Actuals upload not found.", 404)
        return upload

    async def delete_forecast(self, upload_id: int) -> None:
        upload = await self.db.get(OperationalForecastUpload, upload_id)
        if not upload:
            raise OperationalError("Target upload not found.", 404)
        actual_count = await self.db.scalar(
            select(func.count(OperationalActualUpload.id)).where(
                OperationalActualUpload.forecast_upload_id == upload_id
            )
        )
        if actual_count:
            raise OperationalError(
                "Delete the daily actual files for this month before deleting its target.", 409
            )
        await self.db.delete(upload)
        await self.db.commit()

    async def delete_forecast_amendment(self, amendment_id: int) -> None:
        amendment = await self.db.get(OperationalForecastAmendment, amendment_id)
        if not amendment or amendment.deleted_at is not None:
            raise OperationalError("Target amendment not found.", 404)
        forecast = await self.db.get(
            OperationalForecastUpload, amendment.forecast_upload_id
        )
        if not forecast:
            raise OperationalError("Target upload not found.", 404)

        amendment.deleted_at = datetime.now(timezone.utc)
        forecast.row_count = max(0, int(forecast.row_count) - int(amendment.row_count))
        forecast.planned_units = max(
            0, float(forecast.planned_units) - float(amendment.planned_units)
        )
        forecast.po_value = max(0, float(forecast.po_value) - float(amendment.po_value))
        forecast.total_budget = max(
            0, float(forecast.total_budget) - float(amendment.total_budget)
        )
        await self.db.commit()

    async def delete_actual(self, upload_id: int) -> None:
        upload = await self.db.get(OperationalActualUpload, upload_id)
        if not upload:
            raise OperationalError("Actuals upload not found.", 404)
        await self.db.delete(upload)
        await self.db.commit()

    @staticmethod
    def template(kind: str, file_format: str) -> tuple[bytes, str, str]:
        headers = FORECAST_HEADERS if kind == "forecast" else ACTUAL_HEADERS
        stem = "Monthly Target" if kind == "forecast" else "operational-actual-template"
        if file_format == "csv":
            stream = io.StringIO()
            csv.writer(stream).writerow(headers)
            return stream.getvalue().encode("utf-8-sig"), f"{stem}.csv", "text/csv; charset=utf-8"
        if file_format == "xlsx":
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Monthly Target" if kind == "forecast" else "Daily Actuals"
            sheet.append(headers)
            sheet.freeze_panes = "A2"
            for cell in sheet[1]:
                cell.font = Font(bold=True)
            for column in sheet.columns:
                letter = column[0].column_letter
                sheet.column_dimensions[letter].width = max(14, len(str(column[0].value)) + 2)
            output = io.BytesIO()
            workbook.save(output)
            return (
                output.getvalue(),
                f"{stem}.xlsx",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        raise OperationalError("Template format must be csv or xlsx.")

    async def _resolve_range(
        self, date_from: date | None, date_to: date | None
    ) -> tuple[date, date] | None:
        latest_actual = await self.db.scalar(select(func.max(OperationalActualUpload.report_date)))
        latest_forecast = await self.db.scalar(
            select(func.max(OperationalForecastUpload.forecast_month))
        )
        if not latest_actual and not latest_forecast:
            return None
        if date_from and date_to and date_from > date_to:
            raise OperationalError("From date cannot be after To date.")
        if not date_from:
            if date_to:
                date_from = _month_start(date_to)
            else:
                available_months = [
                    value
                    for value in (
                        _month_start(latest_actual) if latest_actual else None,
                        latest_forecast,
                    )
                    if value is not None
                ]
                date_from = max(available_months)
        if not date_to:
            date_to = (
                latest_actual
                if latest_actual and _month_start(latest_actual) == _month_start(date_from)
                else _month_end(date_from)
            )
        if date_from > date_to:
            raise OperationalError("From date cannot be after To date.")
        return date_from, date_to

    async def dashboard(
        self, date_from: date | None = None, date_to: date | None = None
    ) -> OperationalDashboard:
        resolved = await self._resolve_range(date_from, date_to)
        if not resolved:
            return _empty_dashboard()
        date_from, date_to = resolved
        first_month, last_month = _month_start(date_from), _month_start(date_to)

        plan_records = (
            await self.db.execute(
                select(OperationalForecastUpload.forecast_month, OperationalForecastRow)
                .join(
                    OperationalForecastRow,
                    OperationalForecastRow.upload_id == OperationalForecastUpload.id,
                )
                .outerjoin(
                    OperationalForecastAmendment,
                    OperationalForecastRow.source_amendment_id
                    == OperationalForecastAmendment.id,
                )
                .where(
                    OperationalForecastUpload.forecast_month >= first_month,
                    OperationalForecastUpload.forecast_month <= last_month,
                    or_(
                        OperationalForecastRow.source_amendment_id.is_(None),
                        OperationalForecastAmendment.deleted_at.is_(None),
                    ),
                )
            )
        ).all()
        if not plan_records:
            empty = _empty_dashboard()
            empty.date_from = date_from
            empty.date_to = date_to
            return empty

        actual_records = (
            await self.db.execute(
                select(OperationalActualUpload.report_date, OperationalActualRow)
                .join(
                    OperationalActualRow,
                    OperationalActualRow.upload_id == OperationalActualUpload.id,
                )
                .where(
                    OperationalActualUpload.report_date >= date_from,
                    OperationalActualUpload.report_date <= date_to,
                )
            )
        ).all()
        actual_by_row_date = {
            (actual.forecast_row_id, report_date): actual
            for report_date, actual in actual_records
        }

        total = _blank_accumulator()
        category_values: dict[str, dict[str, Any]] = defaultdict(_blank_accumulator)
        asin_values: dict[tuple[str, str], dict[str, Any]] = defaultdict(_blank_accumulator)
        asin_names: dict[tuple[str, str], str] = {}
        plan_by_date: dict[date, list[OperationalForecastRow]] = defaultdict(list)
        plan_by_category_date: dict[
            str, dict[date, list[OperationalForecastRow]]
        ] = defaultdict(lambda: defaultdict(list))
        plan_by_asin_date: dict[
            tuple[str, str], dict[date, list[OperationalForecastRow]]
        ] = defaultdict(lambda: defaultdict(list))

        for month, row in plan_records:
            overlap_start = max(date_from, month, row.effective_from or month)
            overlap_end = min(date_to, _month_end(month))
            if overlap_start > overlap_end:
                continue
            period_dates = list(_dates(overlap_start, overlap_end))
            month_days = _days_in_month(month)
            values = _blank_accumulator()
            values["planned_units"] = float(row.daily_run_rate) * len(period_dates)
            # PO Value is already the validated daily target. Keeping the stored value avoids
            # losing precision when ORM fields with display-scale decimals are loaded.
            values["planned_po_value"] = float(row.po_value) * len(period_dates)
            for component in COMPONENTS:
                monthly = _forecast_component(row, component)
                values["planned"][component] = monthly / month_days * len(period_dates)

            for report_date in period_dates:
                plan_by_date[report_date].append(row)
                plan_by_category_date[row.category][report_date].append(row)
                plan_by_asin_date[(row.category, row.asin)][report_date].append(row)
                actual = actual_by_row_date.get((row.id, report_date))
                if not actual:
                    continue
                actual_units = float(actual.actual_drr)
                values["actual_units"] += actual_units
                values["actual_po_value"] += actual_units * float(actual.po_price)
                monthly_units = float(row.daily_run_rate) * month_days
                for component in COMPONENTS:
                    actual_component = _actual_component(actual, component)
                    monthly_component = _forecast_component(row, component)
                    values["actual"][component] += actual_component
                    if monthly_units:
                        values["adjusted"][component] += (
                            monthly_component / monthly_units * actual_units
                        )

            _merge(total, values)
            _merge(category_values[row.category], values)
            key = (row.category, row.asin)
            _merge(asin_values[key], values)
            asin_names[key] = row.short_name

        categories: list[CategoryOperationalMetrics] = []
        for category, values in category_values.items():
            asin_rows = [
                AsinOperationalMetrics(
                    asin=asin,
                    short_name=asin_names[(cat, asin)],
                    category=cat,
                    timeline=_build_timeline(
                        plan_by_asin_date[(cat, asin)], actual_by_row_date
                    ),
                    **_metrics(asin_values[(cat, asin)]).model_dump(),
                )
                for cat, asin in asin_values
                if cat == category
            ]
            asin_rows.sort(key=lambda item: item.actual_po_value, reverse=True)
            categories.append(
                CategoryOperationalMetrics(
                    category=category,
                    asins=asin_rows,
                    timeline=_build_timeline(
                        plan_by_category_date[category], actual_by_row_date
                    ),
                    **_metrics(values).model_dump(),
                )
            )
        categories.sort(key=lambda item: item.actual_po_value, reverse=True)

        return OperationalDashboard(
            date_from=date_from,
            date_to=date_to,
            covered_days=len(plan_by_date),
            summary=_metrics(total),
            categories=categories,
            timeline=_build_timeline(plan_by_date, actual_by_row_date),
        )

    async def dashboard_export(
        self, date_from: date | None = None, date_to: date | None = None
    ) -> tuple[bytes, str]:
        dashboard = await self.dashboard(date_from=date_from, date_to=date_to)
        stream = io.StringIO()
        writer = csv.writer(stream)
        writer.writerow(
            [
                "Level",
                "Category",
                "ASIN",
                "Short Name",
                "Target Units",
                "Actual Units",
                "Unit Achievement %",
                "Planned PO Value",
                "Actual PO Value",
                "PO Value Variance",
                "Target Spend",
                "Actual Spend",
                "Volume Adjusted Budget",
                "Adjusted Spend Variance",
                "Planned Cost / Unit",
                "Actual Cost / Unit",
                "Target CAC (Target CCOGS + Ads / Target Order)",
                "Actual CAC (CCOGS + Ads / Actual Order)",
                "Planned Spend Utilization %",
                "Actual Spend Utilization %",
                "Contribution Value",
                "Contribution Margin %",
                "Status",
            ]
        )

        def write_row(level: str, category: str, asin: str, short_name: str, row) -> None:
            writer.writerow(
                [
                    level,
                    category,
                    asin,
                    short_name,
                    row.planned_units,
                    row.actual_units,
                    row.unit_achievement_pct,
                    row.planned_po_value,
                    row.actual_po_value,
                    row.po_value_variance,
                    row.planned_spend,
                    row.actual_spend,
                    row.volume_adjusted_budget,
                    row.adjusted_spend_variance,
                    row.planned_cost_per_unit,
                    row.actual_cost_per_unit,
                    row.target_cac,
                    row.cac,
                    row.planned_spend_utilization_pct,
                    row.actual_spend_utilization_pct,
                    row.contribution_value,
                    row.contribution_margin_pct,
                    row.status,
                ]
            )

        write_row("Grand Total", "All Categories", "", "", dashboard.summary)
        for category in dashboard.categories:
            write_row("Category", category.category, "", "", category)
            for asin in category.asins:
                write_row("ASIN", category.category, asin.asin, asin.short_name, asin)
        start = dashboard.date_from.isoformat() if dashboard.date_from else "no-data"
        end = dashboard.date_to.isoformat() if dashboard.date_to else "no-data"
        return stream.getvalue().encode("utf-8-sig"), f"target-vs-achieved-{start}-to-{end}.csv"
